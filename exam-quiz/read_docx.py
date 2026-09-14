from docx import Document
import os

docx_path = r'D:\应用程序开发\刷题\题目和答案pdf\公式截图.docx'

doc = Document(docx_path)

print(f"段落数: {len(doc.paragraphs)}")
print(f"表格数: {len(doc.tables)}")
print(f"内嵌图片数: {len(doc.inline_shapes)}")

print("\n=== 段落内容 ===")
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f"[{i}] {para.text[:200]}")

print("\n=== 表格内容 ===")
for i, table in enumerate(doc.tables):
    print(f"表格{i}: {len(table.rows)}行 x {len(table.columns)}列")
    for j, row in enumerate(table.rows):
        for k, cell in enumerate(row.cells):
            if cell.text.strip():
                print(f"  [{j},{k}] {cell.text[:100]}")
